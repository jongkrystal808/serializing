import inspect
import json
import os
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import ValidationError


BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from app.core.config import CustomerFileConfig, settings  # noqa: E402
from app.core.customers import Customer  # noqa: E402
from app.core.dependencies import get_print_notice_service  # noqa: E402
from app.main import app, lifespan  # noqa: E402
from app.schemas.export import ExportRequest, MAX_EXPORT_ROWS  # noqa: E402
from app.schemas.history import HistoryUpsertRequest, MAX_HISTORY_INCREMENT  # noqa: E402
from app.schemas.sn import GenerateSnRequest, MAX_SN_QTY  # noqa: E402
from app.services.excel_service import ExcelService  # noqa: E402
from app.services.history_service import HistoryService  # noqa: E402
from app.services.print_notice_service import PrintNoticeService  # noqa: E402
from app.services.sn_service import SnService  # noqa: E402


class BackendDesignRegressionTests(unittest.TestCase):
    def test_excel_integer_float_is_rendered_without_decimal_suffix(self) -> None:
        service = ExcelService()
        self.assertEqual(service._apply_parse_rules(50.0, ["trim"]), "50")
        self.assertEqual(service._apply_parse_rules(50.5, ["trim"]), "50.5")
        self.assertEqual(service._apply_parse_rules(0.0, ["trim"]), "0")

    def test_default_excel_workbook_cache_reuses_path_and_invalidates_on_mtime(self) -> None:
        with TemporaryDirectory() as folder:
            path = Path(folder) / "shared.xlsx"
            workbook = Workbook()
            workbook.active.title = "營邦出貨"
            workbook.active.append(["工單"])
            workbook.active.append(["WO-1"])
            workbook.create_sheet("倫飛出貨").append(["MO"])
            workbook["倫飛出貨"].append(["MO-1"])
            workbook.save(path)
            workbook.close()

            original = settings.default_excel
            object.__setattr__(settings, "default_excel", {
                "yingbang": CustomerFileConfig(str(path), "營邦出貨", ["trim"]),
                "lunfei": CustomerFileConfig(str(path), "倫飛出貨", ["trim"]),
            })
            service = ExcelService()
            try:
                with patch.object(service, "_load_workbook", wraps=service._load_workbook) as loader:
                    service.load_from_default_path("yingbang")
                    service.load_from_default_path("lunfei")
                    self.assertEqual(loader.call_count, 1)
                    current = path.stat().st_mtime_ns
                    os.utime(path, ns=(current + 1_000_000_000, current + 1_000_000_000))
                    service.load_from_default_path("yingbang")
                    self.assertEqual(loader.call_count, 2)
            finally:
                object.__setattr__(settings, "default_excel", original)

    def test_generate_sn_schema_rejects_unbounded_quantity(self) -> None:
        with self.assertRaises(ValidationError):
            GenerateSnRequest(customer="yingbang", key="WO", qty=MAX_SN_QTY + 1)
        with self.assertRaises(ValidationError):
            GenerateSnRequest(customer="yingbang", key="WO", qty=0)

    def test_write_schemas_reject_oversized_payload_parts(self) -> None:
        with self.assertRaises(ValidationError):
            HistoryUpsertRequest(customer="yingbang", key="WO", increment=MAX_HISTORY_INCREMENT + 1)
        with self.assertRaises(ValidationError):
            HistoryUpsertRequest(customer="yingbang", key="WO", record="x" * 2_001)
        with self.assertRaises(ValidationError):
            ExportRequest(customer="yingbang", sn_rows=[{}] * (MAX_EXPORT_ROWS + 1))
        with self.assertRaises(ValidationError):
            ExportRequest(customer="yingbang", sn_rows=[{"SN": "x" * 501}])
        with self.assertRaises(ValidationError):
            GenerateSnRequest(
                customer="bng",
                key="WO",
                qty=1,
                provided_serials=["x" * 501],
            )

    def test_lunfei_record_uses_same_week_key_as_serial_counter(self) -> None:
        with TemporaryDirectory() as folder:
            history_service = HistoryService(str(Path(folder) / "lunfei.db"))
            history_service.initialize()
            service = SnService(history_service)
            result = service.generate(
                GenerateSnRequest(
                    customer="lunfei",
                    key="MO001",
                    week_key="2026-W11",
                    qty=1,
                    record="2026-03-12-MO001",
                )
            )
            records = history_service.list_generation_records("lunfei")

        self.assertEqual(result.key, "MO001")
        self.assertEqual(records[0].key, "2026-W11")

    def test_history_database_enables_wal(self) -> None:
        with TemporaryDirectory() as folder:
            service = HistoryService(str(Path(folder) / "wal.db"))
            service.initialize()
            connection = service._connect()
            try:
                journal_mode = connection.execute("PRAGMA journal_mode").fetchone()[0]
            finally:
                connection.close()
        self.assertEqual(str(journal_mode).lower(), "wal")

    def test_http_requests_emit_structured_log(self) -> None:
        with self.assertLogs("app.request", level="INFO") as captured:
            response = TestClient(app).get("/api/health")
        payload = json.loads(captured.records[-1].getMessage())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(payload["event"], "http_request_completed")
        self.assertEqual(payload["method"], "GET")
        self.assertEqual(payload["path"], "/api/health")
        self.assertIn("request_id", payload)
        self.assertIn("duration_ms", payload)

    def test_history_writes_use_sqlite_transaction_not_process_lock(self) -> None:
        source = inspect.getsource(HistoryService)
        self.assertIn('conn.execute("BEGIN IMMEDIATE")', source)
        self.assertNotIn("threading.Lock", source)
        self.assertNotIn("conn.commit()", source)

    def test_history_increment_supports_sqlite_without_returning(self) -> None:
        """歷史遞增需支援舊版 SQLite，並依賴 BEGIN IMMEDIATE 保持原子性。"""
        source = inspect.getsource(HistoryService.upsert_entry).upper()
        self.assertNotIn("RETURNING", source)

    def test_importing_service_modules_does_not_create_database(self) -> None:
        """匯入 service 或 app.main 不得產生 SQLite 檔案或建表。"""
        with TemporaryDirectory() as folder:
            db_path = Path(folder) / "import-side-effect.db"
            environment = os.environ.copy()
            environment["DB_PATH"] = str(db_path)
            environment["PYTHONPATH"] = str(BACKEND_ROOT)
            result = subprocess.run(
                [
                    sys.executable,
                    "-c",
                    (
                        "import app.services.history_service; "
                        "import app.services.print_notice_service; "
                        "import app.main"
                    ),
                ],
                cwd=folder,
                env=environment,
                capture_output=True,
                text=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertFalse(db_path.exists())

    def test_lifespan_initializes_app_scoped_services(self) -> None:
        """DB 建表只在 FastAPI lifespan 啟動時執行。"""
        with TemporaryDirectory() as folder:
            db_path = Path(folder) / "lifespan.db"
            original_db_path = settings.db_path
            object.__setattr__(settings, "db_path", str(db_path))
            test_app = FastAPI(lifespan=lifespan)
            try:
                self.assertFalse(db_path.exists())
                with TestClient(test_app):
                    self.assertTrue(db_path.exists())
                    self.assertIs(
                        test_app.state.sn_service.history_service,
                        test_app.state.history_service,
                    )
                    self.assertIsNotNone(test_app.state.print_notice_service)
            finally:
                object.__setattr__(settings, "db_path", original_db_path)

    def test_openapi_uses_concrete_generic_response_data(self) -> None:
        """各 JSON API 的 data schema 應指向具體模型，而非任意 object。"""
        schemas = app.openapi()["components"]["schemas"]
        expected_response_schemas = {
            "ApiResponse_HealthData_",
            "ApiResponse_ParseExcelResponse_",
            "ApiResponse_GenerateSnResponse_",
            "ApiResponse_HistoryData_",
            "ApiResponse_HistoryUpsertData_",
            "ApiResponse_HistoryResetData_",
            "ApiResponse_PrintNoticeListData_",
            "ApiResponse_PrintNoticeUpsertData_",
            "ApiResponse_PrintNoticeDeleteData_",
        }
        self.assertTrue(expected_response_schemas.issubset(schemas))
        for schema_name in expected_response_schemas:
            data_schema = schemas[schema_name]["properties"]["data"]
            self.assertIn("$ref", data_schema, schema_name)
            self.assertNotIn("additionalProperties", data_schema, schema_name)

        self.assertEqual(
            set(schemas["Customer"]["enum"]),
            {customer.value for customer in Customer},
        )

    def test_print_notice_delete_uses_delete_method_and_dependency(self) -> None:
        with TemporaryDirectory() as folder:
            service = PrintNoticeService(str(Path(folder) / "print-notice.db"))
            service.initialize()
            app.dependency_overrides[get_print_notice_service] = lambda: service
            try:
                client = TestClient(app)
                created = client.post(
                    "/api/print-notice/upsert",
                    json={
                        "customer": "yingbang",
                        "customer_label": "營邦",
                        "workorder_label": "工單",
                        "workorder_value": "WO-A11Y",
                    },
                )
                deleted = client.request(
                    "DELETE",
                    "/api/print-notice",
                    json={"customer": "yingbang", "workorder_value": "WO-A11Y"},
                )
                old_endpoint = client.post(
                    "/api/print-notice/delete",
                    json={"customer": "yingbang", "workorder_value": "WO-A11Y"},
                )
            finally:
                app.dependency_overrides.clear()

        self.assertEqual(created.status_code, 200)
        self.assertEqual(deleted.status_code, 200)
        self.assertTrue(deleted.json()["data"]["removed"])
        self.assertEqual(old_endpoint.status_code, 404)


if __name__ == "__main__":
    unittest.main()
