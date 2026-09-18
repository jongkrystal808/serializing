from __future__ import annotations

import json
import os
import sys
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Callable, List

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
BACKEND_ROOT = PROJECT_ROOT / "backend"
sys.path.insert(0, str(BACKEND_ROOT))

TEST_DB_PATH = BACKEND_ROOT / "data" / "sn_generator_t27.db"


@dataclass
class CaseResult:
    case_id: str
    title: str
    passed: bool
    details: str


def assert_true(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def parse_success_json(response, expected_status: int = 200) -> dict:
    assert_true(response.status_code == expected_status, f"HTTP 狀態不符，預期 {expected_status}，實際 {response.status_code}")
    payload = response.json()
    assert_true(payload.get("success") is True, f"API success 應為 true，實際 {payload.get('success')}")
    assert_true(isinstance(payload.get("data"), dict), "API data 欄位應為 object")
    return payload["data"]


def parse_error_json(response, expected_status: int = 400, expected_code: str | None = None) -> dict:
    assert_true(response.status_code == expected_status, f"HTTP 狀態不符，預期 {expected_status}，實際 {response.status_code}")
    payload = response.json()
    assert_true(payload.get("success") is False, f"API success 應為 false，實際 {payload.get('success')}")
    error = payload.get("error") or {}
    if expected_code:
        assert_true(error.get("code") == expected_code, f"error.code 不符，預期 {expected_code}，實際 {error.get('code')}")
    return payload


def run_case(case_id: str, title: str, fn: Callable[[], str], results: List[CaseResult]) -> None:
    try:
        details = fn()
        results.append(CaseResult(case_id=case_id, title=title, passed=True, details=details))
    except Exception as error:  # noqa: BLE001
        results.append(CaseResult(case_id=case_id, title=title, passed=False, details=str(error)))


def setup_test_db(db_path: Path) -> None:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    if db_path.exists():
        db_path.unlink()
    os.environ["DB_PATH"] = str(db_path)


def build_sample_excel_bytes() -> bytes:
    """建立 T27 所需的最小多工作表 Excel，避免依賴 repository 外部檔案。"""
    workbook = Workbook()
    workbook.remove(workbook.active)
    fixtures = {
        "營邦出貨": (
            ["工單號", "採單號碼", "DDC料號", "品名", "料號", "數量"],
            ["WO001", "PO001", "DDC001", "Model-Y", "PN001", 2],
        ),
        "倫飛出貨": (
            ["MO", "Model", "工單", "P/N", "加工WO#", "對應PCBA", "Q'ty"],
            ["MO001", "BAG017-AAA", "LWO001", "PNL001", "PRO001", "PCBA001", 2],
        ),
        "超恩出貨": (
            ["MO", "機種名稱", "機種料號", "生產數量", "序號區間"],
            ["BMO001", "ModelA", "PART001", 2, "1001~1002"],
        ),
        "KOYA出貨": (
            ["工單", "機種", "PN", "小張貼紙", "滿箱數量"],
            ["CWO001", "KModel", "KPN001", 2, 1],
        ),
    }
    for sheet_name, (headers, row) in fixtures.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def main(db_path: Path | None = None) -> int:
    active_db_path = db_path or TEST_DB_PATH
    original_db_path_env = os.environ.get("DB_PATH")
    setup_test_db(active_db_path)

    from app.core.config import settings  # noqa: WPS433,E402
    from app.main import app  # noqa: WPS433,E402

    original_settings_db_path = settings.db_path
    object.__setattr__(settings, "db_path", str(active_db_path))
    client = TestClient(app)
    client.__enter__()
    sample_excel_name = "t27-generated-fixture.xlsx"
    excel_bytes = build_sample_excel_bytes()

    def parse_excel(customer: str, sheet_name: str, parse_rules: str) -> dict:
        response = client.post(
            "/api/excel/parse",
            data={
                "customer": customer,
                "sheet_name": sheet_name,
                "parse_rules": parse_rules,
            },
            files={
                "file": (
                    sample_excel_name,
                    excel_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        return parse_success_json(response)

    def case_yingbang_flow() -> str:
        parsed = parse_excel("yingbang", "營邦出貨", "arrow")
        assert_true(parsed["rows_count"] >= 1, "營邦解析結果 rows_count 應 >= 1")

        generated_data = parse_success_json(
            client.post(
                "/api/sn/generate",
                json={
                    "customer": "yingbang",
                    "key": "WO001",
                    "qty": 2,
                    "purchase_order": "PO001",
                },
            )
        )
        assert_true(generated_data["generated_count"] == 2, "營邦 generated_count 應為 2")
        assert_true(generated_data["sn_list"][0] == "PO00110001", "營邦第一筆 SN 不符預期")
        assert_true(generated_data["sn_list"][1] == "PO00110002", "營邦第二筆 SN 不符預期")

        parse_success_json(
            client.post(
                "/api/history/upsert",
                json={
                    "customer": "yingbang",
                    "key": "WO001",
                    "increment": 0,
                    "record": "2026-03-12-WO001",
                },
            )
        )

        export_response = client.post(
            "/api/export",
            json={
                "customer": "yingbang",
                "sn_rows": [
                    {"SN": generated_data["sn_list"][0], "Datecode": "D2611", "PN": "PN001"},
                    {"SN": generated_data["sn_list"][1], "Datecode": "D2611", "PN": "PN001"},
                ],
                "file_name": "營邦-SN.xls",
            },
        )
        assert_true(export_response.status_code == 200, "營邦匯出應成功")
        assert_true(len(export_response.content) > 0, "營邦匯出內容不可為空")

        history_data = parse_success_json(client.get("/api/history/yingbang"))
        entries = history_data["entries"]
        matched = next((item for item in entries if item["key"] == "WO001"), None)
        assert_true(bool(matched), "營邦歷史缺少 WO001")
        assert_true(int(matched["last_serial"]) == 2, "營邦歷史流水號應為 2")

        return "營邦：解析/生成/匯出/歷史查詢通過（key=工單 WO001）"

    def case_lunfei_flow_with_week_reset() -> str:
        parsed = parse_excel("lunfei", "倫飛出貨", "arrow")
        assert_true(parsed["rows_count"] >= 1, "倫飛解析結果 rows_count 應 >= 1")

        first = parse_success_json(
            client.post(
                "/api/sn/generate",
                json={
                    "customer": "lunfei",
                    "key": "2026-W11",
                    "qty": 2,
                    "week_key": "2026-W11",
                },
            )
        )
        second = parse_success_json(
            client.post(
                "/api/sn/generate",
                json={
                    "customer": "lunfei",
                    "key": "2026-W11",
                    "qty": 1,
                    "week_key": "2026-W11",
                },
            )
        )
        third = parse_success_json(
            client.post(
                "/api/sn/generate",
                json={
                    "customer": "lunfei",
                    "key": "2026-W12",
                    "qty": 1,
                    "week_key": "2026-W12",
                },
            )
        )
        assert_true(first["previous_serial"] == 0 and first["current_serial"] == 2, "倫飛同週第一次生成結果不符")
        assert_true(second["previous_serial"] == 2 and second["current_serial"] == 3, "倫飛同週接續結果不符")
        assert_true(third["previous_serial"] == 0 and third["current_serial"] == 1, "倫飛跨週重置結果不符")

        parse_success_json(
            client.post(
                "/api/history/upsert",
                json={
                    "customer": "lunfei",
                    "key": "MO001",
                    "increment": 0,
                    "record": "2026-03-12-MO001",
                },
            )
        )

        export_response = client.post(
            "/api/export",
            json={
                "customer": "lunfei",
                "sn_rows": [{"SN": third["sn_list"][0]}],
                "box_row": {
                    "P/N": "PNL001",
                    "加工WO#": "PRO001",
                    "對應PCBA": "PCBA001",
                    "工單": "LWO001",
                    "Model": "BAG017-AAA",
                    "日期": "2026/03/12",
                },
                "file_name": "倫飛-SN.xls",
            },
        )
        assert_true(export_response.status_code == 200, "倫飛匯出應成功")
        assert_true(len(export_response.content) > 0, "倫飛匯出內容不可為空")

        history_data = parse_success_json(client.get("/api/history/lunfei"))
        keys = {item["key"]: int(item["last_serial"]) for item in history_data["entries"]}
        assert_true(keys.get("2026-W11") == 3, "倫飛 2026-W11 歷史流水號應為 3")
        assert_true(keys.get("2026-W12") == 1, "倫飛 2026-W12 歷史流水號應為 1")

        return "倫飛：同週接續與跨週重置通過（W11->W12）"

    def case_bng_flow_with_count_validation() -> str:
        parsed = parse_excel("bng", "超恩出貨", "trim")
        assert_true(parsed["rows_count"] >= 1, "超恩解析結果 rows_count 應 >= 1")

        generated_data = parse_success_json(
            client.post(
                "/api/sn/generate",
                json={
                    "customer": "bng",
                    "key": "BMO001",
                    "qty": 2,
                    "provided_serials": ["1001", "1002"],
                    "record": "2026-03-12-BMO001",
                },
            )
        )
        assert_true(generated_data["generated_count"] == 2, "超恩 generated_count 應為 2")

        export_response = client.post(
            "/api/export",
            json={
                "customer": "bng",
                "sn_rows": [
                    {"序號": "1001", "MAC Address": "001122334455", "uuid1": "無", "uuid2": "", "機種名稱": "ModelA", "BIOS": "BIOS1", "FW": "FW1"},
                    {"序號": "1002", "MAC Address": "001122334456", "uuid1": "無", "uuid2": "", "機種名稱": "ModelA", "BIOS": "BIOS1", "FW": "FW1"},
                ],
                "box_row": {
                    "PO": "BWO001",
                    "Model": "ModelA",
                    "料號": "PART001",
                    "SN": "1001~1002",
                    "思創PN": "PART001",
                    "Date": "2026/03/12",
                },
                "file_name": "超恩-SN.xls",
            },
        )
        assert_true(export_response.status_code == 200, "超恩匯出應成功")
        assert_true(len(export_response.content) > 0, "超恩匯出內容不可為空")

        history_data = parse_success_json(client.get("/api/history/bng"))
        matched = next((item for item in history_data["entries"] if item["key"] == "BMO001"), None)
        assert_true(bool(matched), "超恩歷史缺少 BMO001")
        assert_true(int(matched["last_serial"]) == 2, "超恩歷史流水號應為 2")

        return "超恩：解析/筆數驗證/匯出/歷史查詢通過（key=MO）"

    def case_bng_deleted_marker_cleanup() -> str:
        bng_workbook = Workbook()
        bng_sheet = bng_workbook.active
        bng_sheet.title = "超恩出貨"
        bng_sheet.append(["MO", "機種名稱", "新版BIOS(以此為主)", "IGN FW版本"])
        bng_sheet.append(["BMO-DEL-1", "舊機種[已刪除]暫存[已刪除]CLG078-019M", "A[已刪除]B[已刪除]BIOS-NEW", "FW-KEEP"])
        bng_buffer = BytesIO()
        bng_workbook.save(bng_buffer)
        bng_bytes = bng_buffer.getvalue()

        parsed = parse_success_json(
            client.post(
                "/api/excel/parse",
                data={
                    "customer": "bng",
                    "sheet_name": "超恩出貨",
                    "parse_rules": "trim",
                },
                files={
                    "file": (
                        "bng-deleted-marker.xlsx",
                        bng_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
        )
        assert_true(parsed["rows_count"] == 1, "超恩刪除標記清理測試應有 1 筆資料")
        row = parsed["rows"][0]
        assert_true(row["機種名稱"] == "CLG078-019M", "超恩機種名稱應保留最後一個 [已刪除] 後內容")
        assert_true(row["新版BIOS(以此為主)"] == "BIOS-NEW", "超恩 BIOS 應保留最後一個 [已刪除] 後內容")
        assert_true("[已刪除]" not in "".join(str(value) for value in row.values()), "超恩預覽不應包含 [已刪除] 標記")
        return "超恩：[已刪除] 欄位清理通過（保留最後標記後內容）"

    def case_chg_flow() -> str:
        parsed = parse_excel("chg", "KOYA出貨", "trim")
        assert_true(parsed["rows_count"] >= 1, "KOYA 解析結果 rows_count 應 >= 1")

        generated_data = parse_success_json(
            client.post(
                "/api/sn/generate",
                json={
                    "customer": "chg",
                    "key": "CWO001",
                    "qty": 2,
                    "provided_serials": ["CWO001", "CWO001"],
                    "record": "2026-03-12-CWO001",
                },
            )
        )
        assert_true(generated_data["generated_count"] == 2, "KOYA generated_count 應為 2")

        export_response = client.post(
            "/api/export",
            json={
                "customer": "chg",
                "sn_rows": [
                    {"工單": "CWO001", "PN": "KPN001"},
                    {"工單": "CWO001", "PN": "KPN001"},
                ],
                "box_row": {
                    "PO": "POC1",
                    "PN": "KPN001",
                    "full PN": "FULL001",
                    "DDC PN": "KModel",
                    "DDC LOT": "CWO001",
                    "QTY": "20",
                    "DATE": "2026/03/12",
                },
                "file_name": "KOYA-SN.xls",
            },
        )
        assert_true(export_response.status_code == 200, "KOYA 匯出應成功")
        assert_true(len(export_response.content) > 0, "KOYA 匯出內容不可為空")

        history_data = parse_success_json(client.get("/api/history/chg"))
        matched = next((item for item in history_data["entries"] if item["key"] == "CWO001"), None)
        assert_true(bool(matched), "KOYA 歷史缺少 CWO001")
        assert_true(int(matched["last_serial"]) == 2, "KOYA 歷史流水號應為 2")

        return "KOYA：Label/Box 匯出與歷史查詢通過"

    def case_hmg_flow() -> str:
        hmg_workbook = Workbook()
        hmg_sheet = hmg_workbook.active
        hmg_sheet.title = "Sheet1"
        hmg_sheet.append(["Model", "PN", "EAN Code", "PCBA/Accessories      機種名"])
        hmg_sheet.append(["HMG015-001E", "HX4-06001", "0094333516606", "Pixhawk2.1 Standard Set"])
        hmg_buffer = BytesIO()
        hmg_workbook.save(hmg_buffer)
        hmg_bytes = hmg_buffer.getvalue()
        parsed = parse_success_json(
            client.post(
                "/api/excel/parse",
                data={
                    "customer": "hmg",
                    "sheet_name": "Sheet1",
                    "parse_rules": "none",
                },
                files={
                    "file": (
                        "hmg-sheet1.xlsx",
                        hmg_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
        )
        assert_true(parsed["rows_count"] >= 1, "赫星解析結果 rows_count 應 >= 1")

        parse_success_json(
            client.post(
                "/api/history/upsert",
                json={
                    "customer": "hmg",
                    "key": "HMG015-001E",
                    "increment": 1,
                    "record": "2026-03-18-HMG015-001E",
                },
            )
        )

        export_response = client.post(
            "/api/export",
            json={
                "customer": "hmg",
                "sn_rows": [
                    {"Model": "HMG015-001E", "PN": "HX4-06001", "EAN Code": "0094333516606", "PCBA": "Pixhawk2.1"},
                ],
                "file_name": "20260318-HMG015-001E.xls",
            },
        )
        assert_true(export_response.status_code == 200, "赫星匯出應成功")
        hmg_workbook = load_workbook(filename=BytesIO(export_response.content), data_only=True)
        assert_true("HEX" in hmg_workbook.sheetnames, "赫星匯出 sheet 名稱應為 HEX")
        assert_true(str(hmg_workbook["HEX"]["A2"].value or "") == "HMG015-001E", "赫星 HEX 內容不符預期")

        history_data = parse_success_json(client.get("/api/history/hmg"))
        matched = next((item for item in history_data["entries"] if item["key"] == "HMG015-001E"), None)
        assert_true(bool(matched), "赫星歷史缺少 HMG015-001E")
        assert_true(int(matched["last_serial"]) == 1, "赫星歷史流水號應為 1")

        return "赫星：解析/HEX 匯出/歷史查詢通過（key=Model）"

    def case_failure_scenarios() -> str:
        response_sheet_not_found = client.post(
            "/api/excel/parse",
            data={"customer": "yingbang", "sheet_name": "不存在的工作表", "parse_rules": "arrow"},
            files={
                "file": (
                    sample_excel_name,
                    excel_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        parse_error_json(response_sheet_not_found, expected_status=400, expected_code="SHEET_NOT_FOUND")

        response_mismatch = client.post(
            "/api/sn/generate",
            json={
                "customer": "bng",
                "key": "BMO001",
                "qty": 2,
                "provided_serials": ["only-one"],
            },
        )
        parse_error_json(response_mismatch, expected_status=400, expected_code="INVALID_PROVIDED_SERIALS")

        response_unsupported = client.get("/api/history/not_exists_customer")
        parse_error_json(response_unsupported, expected_status=400, expected_code="UNSUPPORTED_CUSTOMER")

        response_invalid_yingbang = client.post(
            "/api/sn/generate",
            json={
                "customer": "yingbang",
                "key": "WO001",
                "qty": 1,
            },
        )
        parse_error_json(response_invalid_yingbang, expected_status=400, expected_code="MISSING_PURCHASE_ORDER")

        response_invalid_excel = client.post(
            "/api/excel/parse",
            data={"customer": "yingbang", "sheet_name": "營邦出貨", "parse_rules": "arrow"},
            files={
                "file": (
                    "invalid.xlsx",
                    b"not-a-valid-zip-content",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        parse_error_json(response_invalid_excel, expected_status=400, expected_code="INVALID_EXCEL_FILE")

        response_clg_export = client.post(
            "/api/export",
            json={
                "customer": "clg",
                "sn_rows": [{"SN": "CP0001"}],
                "file_name": "Cubepilot-SN.xls",
            },
        )
        assert_true(response_clg_export.status_code == 200, "Cubepilot 匯出應成功")
        clg_workbook = load_workbook(filename=BytesIO(response_clg_export.content), data_only=True)
        assert_true("MES" in clg_workbook.sheetnames, "Cubepilot 匯出 sheet 名稱應為 MES")
        assert_true(str(clg_workbook["MES"]["A2"].value or "") == "CP0001", "Cubepilot MES 內容不符預期")

        return "失敗情境：sheet 不存在 / 筆數不符 / customer 非法 / 缺 purchase_order / 非法 Excel 檔案攔截，且 clg 匯出為 MES"

    results: List[CaseResult] = []
    run_case("T27-01", "營邦完整流程（解析/生成/匯出/歷史）", case_yingbang_flow, results)
    run_case("T27-02", "倫飛完整流程（含週重置）", case_lunfei_flow_with_week_reset, results)
    run_case("T27-03", "超恩完整流程（區間展開與筆數驗證）", case_bng_flow_with_count_validation, results)
    run_case("T27-07", "超恩 [已刪除] 清理", case_bng_deleted_marker_cleanup, results)
    run_case("T27-04", "KOYA 完整流程（Label/Box 匯出）", case_chg_flow, results)
    run_case("T27-05", "赫星完整流程（HEX 匯出）", case_hmg_flow, results)
    run_case("T27-06", "失敗情境測試", case_failure_scenarios, results)

    passed_count = len([item for item in results if item.passed])
    failed_count = len(results) - passed_count
    summary = {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "db_path": str(active_db_path),
        "passed": passed_count,
        "failed": failed_count,
        "results": [
            {
                "case_id": item.case_id,
                "title": item.title,
                "status": "PASS" if item.passed else "FAIL",
                "details": item.details,
            }
            for item in results
        ],
    }
    client.__exit__(None, None, None)
    object.__setattr__(settings, "db_path", original_settings_db_path)
    if original_db_path_env is None:
        os.environ.pop("DB_PATH", None)
    else:
        os.environ["DB_PATH"] = original_db_path_env
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if failed_count == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
