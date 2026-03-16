from io import BytesIO
from typing import Dict, List, Tuple

from openpyxl import Workbook

from app.core.config import settings
from app.core.errors import AppError
from app.schemas.export import ExportRequest

EXPORT_HEADERS: Dict[str, Dict[str, List[str]]] = {
    "yingbang": {
        "SN": ["SN", "Datecode", "PN"],
    },
    "lunfei": {
        "SN": ["SN"],
        "box": ["P/N", "加工WO#", "對應PCBA", "工單", "Model", "日期"],
    },
    "bng": {
        "SN": ["序號", "MAC Address", "uuid1", "uuid2", "機種名稱", "BIOS", "FW"],
        "BOX": ["PO", "Model", "料號", "SN", "思創PN", "Date"],
    },
    "chg": {
        "SN": ["工單", "PN"],
        "BOX": ["PO", "PN", "full PN", "DDC PN", "DDC LOT", "QTY", "DATE"],
    },
}


class ExportService:
    def export(self, payload: ExportRequest) -> Tuple[str, bytes, str]:
        customer = payload.customer.strip()
        if customer not in settings.allowed_customers:
            raise AppError(
                f"不支援的 customer：{customer}",
                code="UNSUPPORTED_CUSTOMER",
                details={"allowed_customers": list(settings.allowed_customers)},
            )
        if not payload.sn_rows:
            raise AppError("sn_rows 不可為空", code="EMPTY_SN_ROWS")

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        config = EXPORT_HEADERS[customer]
        self._append_sn_sheet(workbook, config["SN"], payload.sn_rows)
        if "box" in config:
            self._append_single_row_sheet(workbook, "box", config["box"], payload.box_row)
        if "BOX" in config:
            self._append_single_row_sheet(workbook, "BOX", config["BOX"], payload.box_row)

        output = BytesIO()
        workbook.save(output)
        content = output.getvalue()
        safe_label = customer if customer else "客戶"
        filename = payload.file_name.strip() or f"{safe_label}-SN.xls"
        return (
            filename,
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @staticmethod
    def _append_sn_sheet(workbook: Workbook, headers: List[str], rows: List[Dict[str, str]]) -> None:
        sheet = workbook.create_sheet("SN")
        sheet.append(headers)
        for row in rows:
            sheet.append([str(row.get(column, "")) for column in headers])

    @staticmethod
    def _append_single_row_sheet(
        workbook: Workbook,
        sheet_name: str,
        headers: List[str],
        row: Dict[str, str],
    ) -> None:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        sheet.append([str(row.get(column, "")) for column in headers])


export_service = ExportService()
