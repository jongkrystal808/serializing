from io import BytesIO
from pathlib import Path
from typing import Dict, List
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.config import settings
from app.core.errors import AppError
from app.schemas.excel import ParseExcelRequest, ParseExcelResponse

ARROW_PATTERN = ("->", "→", ">")
OLE2_SIGNATURE = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"
ZIP_SIGNATURE = b"PK\x03\x04"
DELETED_MARKER = "[已刪除]"

COLUMN_ALIASES: Dict[str, Dict[str, List[str]]] = {
    "yingbang": {
        "WORK_ORDER": ["工單號", "工單", "MO", "mo", "製令", "工單編號"],
        "PURCHASE_ORDER": ["採單號碼", "採單", "採購單號", "PO", "PO#"],
        "DDC_PART_NO": ["DDC料號", "DDC 料號", "DDC PN"],
        "PRODUCT_NAME": ["品名", "產品名稱", "MODEL", "Model", "型號"],
        "PN": ["料號", "P/N", "PN", "part number"],
        "QTY": ["數量", "Q'ty", "QTY", "Qty"],
    },
    "lunfei": {
        "MO": ["MO", "mo"],
        "MODEL": ["Model", "MODEL", "model"],
        "WORK_ORDER": ["工單", "工單號", "工單編號"],
        "PN": ["P/N", "PN", "料號"],
        "PROCESS_WO": ["加工WO#", "加工WO", "加工工單"],
        "PCBA": ["對應PCBA", "PCBA", "對應板號"],
        "QTY": ["Q'ty", "QTY", "Qty", "數量"],
    },
    "bng": {
        "MO": ["MO", "mo"],
        "DATE": ["日期", "Date"],
        "WORK_ORDER": ["工單", "工單號"],
        "MODEL": ["機種名稱", "Model", "MODEL"],
        "PART_NO": ["機種料號", "料號", "P/N", "PN"],
        "QTY": ["生產數量", "數量", "QTY", "Qty"],
        "MAC_RANGE": ["MAC  Address", "MAC Address", "MAC區間"],
        "MAC_QTY": ["MAC數量"],
        "MAC_BOARD_QTY": ["MAC板子用量數量", "MAC板用量數量"],
        "SN_RANGE": ["序號區間", "SN區間"],
        "UUID_RANGE": ["UUID區間", "UUID"],
        "BIOS": ["新版BIOS(以此為主)", "新版BIOS", "BIOS"],
        "FW": ["IGN FW版本", "FW", "FW版本"],
    },
    "chg": {
        "WORK_ORDER": ["工單", "工單號"],
        "QTY": ["小張貼紙", "QTY", "Qty", "數量"],
        "MODEL": ["機種", "Model", "MODEL"],
        "PN": ["PN", "P/N", "料號"],
        "BATCH": ["批量"],
        "PO": ["PO", "PO#"],
        "BOX_QTY": ["滿箱數量", "QTY", "Qty"],
        "DEMAND": ["需求"],
        "TAIL_QTY": ["尾數數量"],
        "FULL_PN": ["full PN", "FULL PN", "Full PN"],
    },
    "hmg": {
        "MODEL": ["Model", "MODEL", "model"],
        "PN": ["PN", "P/N", "PK", "料號"],
        "EAN": ["EAN Code", "EAN", "條碼"],
        "PCBA": ["PCBA/Accessories      機種名", "PCBA/Accessories 機種名", "PCB/AAccessories", "PCBA"],
    },
    "clg": {
        "MODEL": ["機種名", "機種", "Model", "MODEL"],
        "QRCODE": ["小張QRCODE", "QRCODE"],
        "CUBE_STICKER": ["Cube測試用貼紙", "Cube測試貼紙"],
        "NOTE": ["note", "Note", "備註"],
    },
}


class ExcelService:
    def load_from_default_path(self, customer: str) -> ParseExcelResponse:
        """依 config 設定的固定路徑直接讀取 Excel，不需使用者上傳。"""
        from pathlib import Path as _Path
        cfg = settings.default_excel.get(customer)
        if cfg is None:
            raise AppError(
                f"customer '{customer}' 尚未設定預設檔案路徑",
                code="NO_DEFAULT_PATH",
                status_code=404,
            )
        file_path = _Path(cfg.path)
        if not file_path.exists():
            raise AppError(
                f"找不到預設 Excel 檔案：{cfg.path}",
                code="DEFAULT_FILE_NOT_FOUND",
                status_code=404,
                details={"path": cfg.path},
            )
        file_bytes = file_path.read_bytes()
        payload = ParseExcelRequest(
            customer=customer,
            sheet_name=cfg.sheet_name,
            file_name=file_path.name,
            parse_rules=cfg.parse_rules,
        )
        return self.parse(payload, file_bytes)

    def parse(self, payload: ParseExcelRequest, file_bytes: bytes) -> ParseExcelResponse:
        customer = payload.customer.strip()
        if customer not in settings.allowed_customers:
            raise AppError(
                f"不支援的 customer：{customer}",
                code="UNSUPPORTED_CUSTOMER",
                details={"allowed_customers": list(settings.allowed_customers)},
            )
        if not file_bytes:
            raise AppError("上傳檔案內容為空", code="EMPTY_FILE")

        workbook_type, workbook = self._load_workbook(payload.file_name, file_bytes)
        try:
            if customer == "hmg":
                parsed_rows, resolved_columns = self._parse_hmg_columnar_rows(
                    workbook_type=workbook_type,
                    workbook=workbook,
                    sheet_name=payload.sheet_name,
                    parse_rules=payload.parse_rules,
                )
                return ParseExcelResponse(
                    customer=customer,
                    sheet_name=payload.sheet_name,
                    file_name=payload.file_name,
                    rows_count=len(parsed_rows),
                    rows=parsed_rows,
                    resolved_columns=resolved_columns,
                )

            sheet_names = self._get_sheet_names(workbook_type, workbook)
            if payload.sheet_name not in sheet_names:
                raise AppError(
                    f"找不到工作表：{payload.sheet_name}",
                    code="SHEET_NOT_FOUND",
                    details={"available_sheets": sheet_names},
                )

            rows_iter = self._iter_rows(workbook_type, workbook, payload.sheet_name)
            headers = self._read_headers(rows_iter)
            if not headers:
                raise AppError("Excel 表頭為空，無法解析", code="EMPTY_HEADER")

            parsed_rows = self._read_rows(rows_iter, headers, payload.parse_rules)
            if customer == "bng":
                parsed_rows = self._sanitize_bng_rows(parsed_rows)
            resolved_columns = self._resolve_columns(customer, headers)
            return ParseExcelResponse(
                customer=customer,
                sheet_name=payload.sheet_name,
                file_name=payload.file_name,
                rows_count=len(parsed_rows),
                rows=parsed_rows,
                resolved_columns=resolved_columns,
            )
        finally:
            self._close_workbook(workbook_type, workbook)

    @staticmethod
    def _resolve_reader_order(file_name: str, file_bytes: bytes) -> List[str]:
        extension = Path(str(file_name or "")).suffix.lower()
        looks_like_xls = extension == ".xls" or file_bytes.startswith(OLE2_SIGNATURE)
        looks_like_xlsx = extension in {".xlsx", ".xlsm", ".xltx", ".xltm"} or file_bytes.startswith(ZIP_SIGNATURE)
        if looks_like_xls and not looks_like_xlsx:
            return ["xls", "xlsx"]
        return ["xlsx", "xls"]

    def _load_workbook(self, file_name: str, file_bytes: bytes):
        errors = {}
        xls_reader_missing = False
        for reader in self._resolve_reader_order(file_name, file_bytes):
            if reader == "xlsx":
                try:
                    return "xlsx", self._load_xlsx_workbook(file_bytes)
                except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
                    errors["xlsx"] = str(error)
                    continue
            try:
                return "xls", self._load_xls_workbook(file_bytes)
            except ImportError:
                xls_reader_missing = True
                errors["xls"] = "xlrd not installed"
            except Exception as error:  # noqa: BLE001
                errors["xls"] = str(error)

        if xls_reader_missing:
            raise AppError(
                "伺服器尚未安裝 .xls 解析元件（xlrd），請聯繫管理員安裝後重試。",
                code="XLS_SUPPORT_NOT_INSTALLED",
                status_code=500,
                details={"required_package": "xlrd"},
            )

        raise AppError(
            "Excel 檔案格式無法解析，請確認為有效的 .xls 或 .xlsx 檔案後重試。",
            code="INVALID_EXCEL_FILE",
            details={"reasons": errors},
        )

    @staticmethod
    def _load_xlsx_workbook(file_bytes: bytes):
        try:
            return load_workbook(
                filename=BytesIO(file_bytes),
                data_only=True,
                read_only=True,
            )
        except (BadZipFile, InvalidFileException, OSError, ValueError):
            raise

    @staticmethod
    def _load_xls_workbook(file_bytes: bytes):
        import xlrd

        return xlrd.open_workbook(file_contents=file_bytes, on_demand=True)

    @staticmethod
    def _get_sheet_names(workbook_type: str, workbook) -> List[str]:
        if workbook_type == "xlsx":
            return list(workbook.sheetnames)
        return list(workbook.sheet_names())

    @staticmethod
    def _iter_rows(workbook_type: str, workbook, sheet_name: str):
        if workbook_type == "xlsx":
            sheet = workbook[sheet_name]
            return sheet.iter_rows(values_only=True)

        sheet = workbook.sheet_by_name(sheet_name)

        def iter_xls_rows():
            for row_index in range(sheet.nrows):
                yield [sheet.cell_value(row_index, col_index) for col_index in range(sheet.ncols)]

        return iter_xls_rows()

    @staticmethod
    def _close_workbook(workbook_type: str, workbook) -> None:
        if workbook_type == "xlsx":
            workbook.close()
            return
        if hasattr(workbook, "release_resources"):
            workbook.release_resources()

    @staticmethod
    def _read_headers(rows_iter) -> List[str]:
        for row in rows_iter:
            candidates = [str(cell or "").strip() for cell in row]
            if any(candidates):
                return candidates
        return []

    def _read_rows(self, rows_iter, headers: List[str], parse_rules: List[str]) -> List[Dict[str, str]]:
        rows: List[Dict[str, str]] = []
        for row in rows_iter:
            parsed = {}
            is_empty = True
            for index, header in enumerate(headers):
                header_text = str(header or "").strip()
                if not header_text:
                    continue
                raw_value = row[index] if index < len(row) else ""
                value = self._apply_parse_rules(raw_value, parse_rules)
                parsed[header_text] = value
                if value:
                    is_empty = False
            if not is_empty:
                rows.append(parsed)
        return rows

    def _apply_parse_rules(self, value, parse_rules: List[str]) -> str:
        text = str(value or "").strip()
        active_rules = parse_rules or ["arrow"]
        for rule in active_rules:
            if rule == "trim":
                text = text.strip()
            if rule == "arrow":
                text = self._parse_arrow(text)
        return text

    @staticmethod
    def _parse_arrow(text: str) -> str:
        value = str(text or "").strip()
        if not value:
            return ""
        for marker in ARROW_PATTERN:
            if marker in value:
                parts = [segment.strip() for segment in value.split(marker) if segment.strip()]
                if parts:
                    value = parts[-1]
        return value

    def _sanitize_bng_rows(self, rows: List[Dict[str, str]]) -> List[Dict[str, str]]:
        sanitized_rows: List[Dict[str, str]] = []
        for row in rows:
            cleaned = {header: self._strip_deleted_prefix(value) for header, value in row.items()}
            if any(cleaned.values()):
                sanitized_rows.append(cleaned)
        return sanitized_rows

    @staticmethod
    def _strip_deleted_prefix(text: str) -> str:
        value = str(text or "").strip()
        marker_index = value.rfind(DELETED_MARKER)
        if marker_index == -1:
            return value
        return value[marker_index + len(DELETED_MARKER) :].strip()

    def _resolve_columns(self, customer: str, headers: List[str]) -> Dict[str, str]:
        aliases = COLUMN_ALIASES.get(customer, {})
        resolved: Dict[str, str] = {}
        for code, candidates in aliases.items():
            resolved_header = self._find_header(headers, candidates)
            if resolved_header:
                resolved[code] = resolved_header
        return resolved

    @staticmethod
    def _find_header(headers: List[str], candidates: List[str]) -> str:
        normalized_headers = {ExcelService._normalize(header): header for header in headers if header}
        for candidate in candidates:
            found = normalized_headers.get(ExcelService._normalize(candidate))
            if found:
                return found
        return ""

    @staticmethod
    def _normalize(text: str) -> str:
        return str(text or "").strip().replace(" ", "").lower()

    def _parse_hmg_columnar_rows(
        self,
        workbook_type: str,
        workbook,
        sheet_name: str,
        parse_rules: List[str],
    ) -> tuple[List[Dict[str, str]], Dict[str, str]]:
        sheet_names = self._get_sheet_names(workbook_type, workbook)
        if sheet_name not in sheet_names:
            raise AppError(
                f"找不到工作表：{sheet_name}",
                code="SHEET_NOT_FOUND",
                details={"available_sheets": sheet_names},
            )

        matrix = self._read_sheet_matrix(workbook_type, workbook, sheet_name)
        if not matrix:
            raise AppError("Excel 內容為空，無法解析赫星資料", code="EMPTY_HMG_SHEET")

        model_rows: List[int] = []
        pn_rows: List[int] = []
        ean_rows: List[int] = []
        pcba_rows: List[int] = []

        for row_index, row in enumerate(matrix):
            label = self._normalize_hmg_label(row[0] if row else "")
            if not label:
                continue
            if label in {"model"}:
                model_rows.append(row_index)
            if label in {"pn", "pk", "p/n"}:
                pn_rows.append(row_index)
            if label in {"eancode", "ean"}:
                ean_rows.append(row_index)
            if label in {
                "pcba/accessories機種名",
                "pcba/accessories",
                "pcbaaccessories機種名",
                "pcbaaccessories",
                "pcb/aaccessories",
                "pcba",
            }:
                pcba_rows.append(row_index)

        if not model_rows:
            raise AppError(
                "赫星解析失敗：找不到 Model 標籤列",
                code="HMG_LAYOUT_INVALID",
            )

        max_cols = max(len(row) for row in matrix)
        parsed_rows: List[Dict[str, str]] = []

        for col in range(1, max_cols):
            model_values = self._collect_hmg_values(matrix, model_rows, col, parse_rules)
            if not model_values:
                continue

            pn_value = self._collect_hmg_first_value(matrix, pn_rows, col, parse_rules)
            ean_value = self._collect_hmg_first_value(matrix, ean_rows, col, parse_rules)
            pcba_value = self._collect_hmg_first_value(matrix, pcba_rows, col, parse_rules)

            for model_value in model_values:
                parsed_rows.append(
                    {
                        "Model": model_value,
                        "PN": pn_value,
                        "EAN Code": ean_value,
                        "PCBA/Accessories      機種名": pcba_value,
                    }
                )

        if not parsed_rows:
            raise AppError(
                "赫星解析失敗：未找到可用的 Model 資料",
                code="HMG_LAYOUT_INVALID",
            )

        resolved_columns = {
            "MODEL": "Model",
            "PN": "PN",
            "EAN": "EAN Code",
            "PCBA": "PCBA/Accessories      機種名",
        }
        return parsed_rows, resolved_columns

    @staticmethod
    def _normalize_hmg_label(value: str) -> str:
        return str(value or "").strip().replace(" ", "").replace("\n", "").lower()

    def _read_sheet_matrix(self, workbook_type: str, workbook, sheet_name: str) -> List[List[str]]:
        matrix: List[List[str]] = []
        rows_iter = self._iter_rows(workbook_type, workbook, sheet_name)
        max_cols = 0
        for row in rows_iter:
            values = [str(cell or "").strip() for cell in row]
            matrix.append(values)
            if len(values) > max_cols:
                max_cols = len(values)
        if max_cols == 0:
            return []
        return [row + [""] * (max_cols - len(row)) for row in matrix]

    def _collect_hmg_values(
        self,
        matrix: List[List[str]],
        row_indexes: List[int],
        col: int,
        parse_rules: List[str],
    ) -> List[str]:
        values: List[str] = []
        seen = set()
        for row_index in row_indexes:
            if row_index < 0 or row_index >= len(matrix):
                continue
            raw = matrix[row_index][col] if col < len(matrix[row_index]) else ""
            value = self._apply_parse_rules(raw, parse_rules)
            if not value:
                continue
            if value in seen:
                continue
            seen.add(value)
            values.append(value)
        return values

    def _collect_hmg_first_value(
        self,
        matrix: List[List[str]],
        row_indexes: List[int],
        col: int,
        parse_rules: List[str],
    ) -> str:
        for row_index in row_indexes:
            if row_index < 0 or row_index >= len(matrix):
                continue
            raw = matrix[row_index][col] if col < len(matrix[row_index]) else ""
            value = self._apply_parse_rules(raw, parse_rules)
            if value:
                return value
        return ""


excel_service = ExcelService()